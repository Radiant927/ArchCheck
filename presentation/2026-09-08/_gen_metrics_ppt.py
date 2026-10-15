# Generate the English "Evaluation Metrics" one-pager (editable pptx).
# White background, black text — plain academic style.
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x44, 0x44, 0x44)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank


def add_text(left, top, width, height, name="Times New Roman"):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for p in (name, "Times New Roman"):
        pass
    return tb, tf


# ---- Title ----
_, tf = add_text(0.4, 0.16, 12.5, 0.6)
p = tf.paragraphs[0]
r = p.add_run(); r.text = "Evaluation Metrics — one row of results per tested model"
r.font.size = Pt(24); r.font.bold = True; r.font.name = "Times New Roman"; r.font.color.rgb = BLACK

# ---- Subtitle ----
_, tf = add_text(0.4, 0.78, 12.5, 0.45)
p = tf.paragraphs[0]
r = p.add_run()
r.text = ("Every IFC file goes through the same pipeline and comes out as one 18-metric row. "
          "Below is a real one: FZK-Haus (Archicad, IFC2X3).")
r.font.size = Pt(13); r.font.name = "Times New Roman"; r.font.color.rgb = GREY


def metric_line(tf, metric, value, desc, blank_before=False):
    p = tf.add_paragraph() if tf.paragraphs[0].text or blank_before else tf.paragraphs[0]
    if blank_before:
        p = tf.add_paragraph()
    p.space_after = Pt(2)
    r = p.add_run(); r.text = f"{metric}  "
    r.font.size = Pt(12); r.font.bold = True; r.font.name = "Courier New"; r.font.color.rgb = BLACK
    r2 = p.add_run(); r2.text = f"{value}  —  "
    r2.font.size = Pt(12); r2.font.bold = True; r2.font.name = "Times New Roman"; r2.font.color.rgb = BLACK
    r3 = p.add_run(); r3.text = desc
    r3.font.size = Pt(11); r3.font.name = "Times New Roman"; r3.font.color.rgb = GREY


def group(left, top, width, title, rows):
    _, tf = add_text(left, top, width, 0.3)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = title
    r.font.size = Pt(13); r.font.bold = True; r.font.name = "Times New Roman"; r.font.color.rgb = BLACK
    _, tf2 = add_text(left, top + 0.3, width, 5.6)
    for (m, v, d) in rows:
        metric_line(tf2, m, v, d)
    return tf2


G1 = [("category", "real-project", "file source"),
      ("file", "FZK-Haus_Archicad11_IFC2x3.ifc", "file name"),
      ("schema", "IFC2X3", "IFC version"),
      ("size_mb", "3.96", "file size (MB)")]

G2 = [("A_robust", "100", "opens, parses, reviews without crashing"),
      ("B_extraction", "100", "share of element classes extracted  ·  weight 25%  (did we read it all?)"),
      ("C_fieldcov", "84.0", "fields filled on judged elements  ·  weight 20%"),
      ("D_accuracy", "100", "dimensions inside real-world ranges  ·  weight 15%"),
      ("E_ruleactivation", "61.3", "share of rules that really judged  ·  weight 15%"),
      ("F_perf", "100", "parse speed, s per MB  ·  weight 5%"),
      ("score", "91.0", "weighted total over the six axes"),
      ("grade", "S", "S≥90 · A≥80 · B≥65 · C≥50 · D<50")]

G3 = [("qa_score", "100", "model health check given by the QA pass"),
      ("storeys", "2", "storeys found"),
      ("rooms", "7", "rooms found"),
      ("issues", "26", "violations actually flagged"),
      ("not_checked", "16", "could not judge — data missing (reported honestly)"),
      ("parse_secs", "5.42", "seconds to parse the file")]

group(0.4, 1.35, 6.15, "①  Which file is this?", G1)
group(0.4, 3.15, 12.5, "②  How well did the platform read & review it?  (six weighted axes → score + grade)", G2)
group(0.4, 5.75, 12.5, "③  What did the review produce?  (the model's own ledger)", G3)

# raw csv row footnote
_, tf = add_text(0.4, 7.05, 12.5, 0.4)
p = tf.paragraphs[0]
r = p.add_run()
r.text = ("Raw CSV row:  real-project · FZK-Haus_Archicad11_IFC2x3.ifc · IFC2X3 · 3.96 · 91.0 · S · "
          "100 · 100 · 84.0 · 100 · 61.3 · 100 · 100 · 2 · 7 · 26 · 16 · 5.42")
r.font.size = Pt(8); r.font.name = "Courier New"; r.font.color.rgb = GREY

# speaker notes
notes = slide.notes_slide.notes_text_frame
notes.text = ("Each tested model = one row of 18 metrics in 3 groups: which file, six weighted "
              "axes plus the weighted total & grade, and what the review produced. "
              "FZK-Haus (Archicad, IFC2X3) scores 91.0 → S. Note weight 25% sits on B_extraction: "
              "the question the platform must answer is 'did we read it all?'. "
              "not_checked=16 means 16 checks honestly reported as unjudgeable for missing data.")

out = r"C:\Users\32800\Desktop\NTU pre\ArchCheck\presentation\2026-09-08\Evaluation_Metrics_FZK-Haus.pptx"
prs.save(out)
print("saved:", out)
