# -*- coding: utf-8 -*-
"""
ArchCheck 汇报 PPT 生成器（V2 重构版 · 9 页）— 纯白底 · 宋体 · 中文
结构：封面 → 上周两个问题/方案 → 自生成IFC → 读取算法链 → 文件/RAG/比对 → 5Agent/四专业 → 平面视图 → 动态阈值 → 下一步
讲解词写入各页演讲者备注。字体统一宋体(SimSun)、字号较上版加大。
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

FONT = "SimSun"          # 宋体
BLACK = RGBColor(0x00, 0x00, 0x00)
GRAY = RGBColor(0x50, 0x50, 0x50)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# 字号（加大版）
TITLE = 30
H = 20      # 小节标题
B = 17      # 正文
S = 15      # 子项

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _ea(run, name=FONT):
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", name)


def blank(p):
    s = p.slides.add_slide(p.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    return s


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def _run(p, text, size, bold, color):
    p.text = text
    run = p.runs[0] if p.runs else p.add_run()
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    _ea(run)


def title(slide, text, tag=None):
    if tag:
        t = slide.shapes.add_textbox(Inches(0.6), Inches(0.38), Inches(12.1), Inches(0.35))
        p = t.text_frame.paragraphs[0]
        _run(p, tag, 13, True, GRAY)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(0.72 if tag else 0.45), Inches(12.1), Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    _run(p, text, TITLE, True, BLACK)


def para(tf, text, size, bold=False, first=False, space_after=6, color=BLACK):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    _run(p, text, size, bold, color)
    p.space_after = Pt(space_after)
    return p


def block(slide, x, y, w, blocks, h_size=H, b_size=B, s_size=S):
    box = slide.shapes.add_textbox(x, y, w, Inches(5.9))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in blocks:
        kind = item[0]
        text = item[1] if len(item) > 1 else ""
        if kind == "h":
            para(tf, text, h_size, bold=True, first=first, space_after=8)
        elif kind == "b":
            para(tf, text, b_size, first=first, space_after=6)
        elif kind == "s":
            para(tf, text, s_size, first=first, space_after=6, color=GRAY)
        elif kind == "blank":
            para(tf, "", s_size, first=first, space_after=4)
        first = False
    return box


def cover(p):
    s = blank(p)
    items = [
        (2.1, "ArchCheck — 基于 AI 的多层建筑规范智能审查平台", 32, True, BLACK),
        (3.3, "从 BIM 解析到动态知识库的技术落地", 20, False, GRAY),
        (5.2, "汇报人：黄韵霖 / YCaiSG / Lihui_PC", 16, False, BLACK),
        (5.7, "2026-08-13", 14, False, GRAY),
    ]
    for y, text, sz, bd, clr in items:
        box = s.shapes.add_textbox(Inches(0.9), Inches(y), Inches(11.5), Inches(0.6))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        _run(p, text, sz, bd, clr)
    notes(s, "各位老师好。我们团队的项目叫 ArchCheck——给建筑做合规体检的智能平台。输入三维 BIM 模型，自动对照规范逐条检查，输出带依据、可追溯的审查报告。今天汇报这周进展，重点讲清数据怎么读、规范怎么管、判定怎么做，以及遇到的真实问题和处理方式。")


def s02(p):
    s = blank(p)
    title(s, "上周的两个问题 → 今天的解决方案")
    block(s, Inches(0.6), Inches(1.7), Inches(5.9), [
        ("h", "上周的两个问题"),
        ("b", "① 数据源：读平面图，不读 BIM"),
        ("s", "之前直接读取 2D 平面图，而非行业通用的三维 BIM 模型"),
        ("s", "2D 图纸信息有限：无房间面积、构件属性、空间关系 → 撑不起专业审查"),
        ("b", "② 判定方式：让通用 LLM 直接判规范"),
        ("s", "之前判定由通用大模型完成，规范依据依赖模型内置/联网知识"),
        ("s", "范围小、易幻觉：实测 FC23 6 处条款号全部指错内容；不可复现、不可审计"),
    ], s_size=S)
    block(s, Inches(7.0), Inches(1.7), Inches(5.7), [
        ("h", "今天的解决方案"),
        ("b", "① 数据端：自生成 IFC + 确定性解析算法链"),
        ("s", "自建 IFC 样例彻底搞懂 BIM 数据结构"),
        ("s", "算法链：单位推断 / 门宽三级回退 / 空间关联 / 疏散距离 / 结构截面"),
        ("b", "② 资料端：规范知识库（RAG）"),
        ("s", "免费规范全文落盘（Fire Code、GB 55037 等 11 份）+ 付费标准索引"),
        ("s", "判定带规范依据 code_ref，可追溯、可审计"),
    ], s_size=S)
    notes(s, "上周老师指出两个方向性问题。第一，数据源：我们之前读的是二维平面图，但行业通用、信息最完整的是三维 BIM 模型，2D 没有房间面积、构件属性、空间关系，做不了专业审查。第二，判定方式：之前让通用大模型直接判规范，范围小、还会编——实测原来规则引用的 Fire Code 条款号 6 处全部指向错误内容，判定还不可复现。所以我们从数据端自己生成 IFC、做扎实读取算法链；资料端把规范整理成知识库。下面逐一展开。")


def s03(p):
    s = blank(p)
    title(s, "为什么自生成 IFC + 生成步骤")
    block(s, Inches(0.6), Inches(1.55), Inches(12.1), [
        ("h", "为什么自建 IFC"),
        ("b", "① 学习 BIM 结构：亲手构造实体/关系/属性（IfcSpace、IfcRelAggregates、Pset_DoorCommon）"),
        ("b", "② 检验阶段「真值回归」：样例真值已知，解析对得上 = 算法正确；快、可控、可复现"),
        ("b", "③ 真实模型难搜集/下载：直链失效（Duplex）、需 GitHub（Schependomlaan）、网络不稳（TLS 失败/连接重置）；Snowdon 142MB 太重"),
        ("blank",),
        ("h", "生成步骤（scripts/generate_office_ifc.py）"),
        ("s", "1. 规划布局：3 层 × 8 房间/层（办公室/会议室/楼梯间/走廊），占地 38m×8.5m"),
        ("s", "2. ifcopenshell.file(schema=IFC4) 建空模型，声明单位 IfcUnitAssignment（米）"),
        ("s", "3. 工程骨架：IfcProject → IfcSite → IfcBuilding → 3×IfcBuildingStorey"),
        ("s", "4. 房间 IfcSpace → IfcRelAggregates 挂楼层（关键坑：解析器靠聚合关系识别房间）"),
        ("s", "5. 构件（墙/门/柱/梁/板）→ IfcRelContainedInSpatialStructure 挂楼层"),
        ("s", "6. 几何：手动 IfcFacetedBrep 面片（绕开 ifcgeom 不支持拉伸几何的限制）"),
        ("s", "7. 门：Pset_DoorCommon.OverallWidth=0.9 + IfcRelSpaceBoundary（喂三级回退第一级）"),
        ("s", "8. 房间类型 Pset_SpaceCommon.Reference；输出 office_3f.ifc（88KB）"),
        ("b", "验证：QA 90；24 房间；面积 95/36/24 m²；门宽 0.9m——与真值完全对上", ),
    ], b_size=16, s_size=14)
    notes(s, "为什么自建 IFC？三条：一是学习成本——亲手构造 IFC，才能真正理解实体、关系、属性三件事；二是检验——自建样例真值是我亲手写的，解析对得上就证明算法对，是最干净的回归测试；三是现实限制——真实模型要么直链失效、要么要上 GitHub、要么 142MB 太重。生成步骤就是「写一份 IFC」：建工程骨架、房间用聚合关系挂楼层、构件用包含关系挂楼层、几何用手写面片、门带标准属性。最关键的一个坑：房间必须用 IfcRelAggregates 挂楼层，用另一种关系解析器就识别不到——这是我们踩过、修好的。")


def s04(p):
    s = blank(p)
    title(s, "读取 IFC 的完整算法链", tag="数据端 · 确定性解析")
    block(s, Inches(0.6), Inches(1.6), Inches(12.1), [
        ("s", "1. 打开模型：ifcopenshell.open(path)"),
        ("s", "2. 单位推断：读 IfcUnitAssignment → 缺失则按楼层标高数值反推（米/毫米）"),
        ("s", "3. 楼层内构件：IfcRelContainedInSpatialStructure（门/窗/墙/柱/梁/板）"),
        ("s", "4. 楼层内空间：IfcRelAggregates（房间 IfcSpace）"),
        ("s", "5. 房间面积：ifcgeom.create_shape 曲面化 → XY 投影 → 凸包 → 鞋带公式"),
        ("s", "6. 门宽三级回退：Pset_DoorCommon.OverallWidth → 名称正则（mm/英寸）→ 几何包围盒"),
        ("s", "7. 门→房间：IfcRelSpaceBoundary 精确关联 → 几何兜底（质心距 ≤0.5m 且同楼层）"),
        ("s", "8. 疏散距离：房间邻接图（共享门连边）+ Dijkstra 最短路"),
        ("s", "9. 结构截面：柱/梁/板截面 + 跨度 span_m"),
        ("blank",),
        ("h", "模型可用性 QA 与输出"),
        ("b", "QA：单位 / 空间完整性 / 几何有效性 / 空间边界 4 项 → 健康度 0-100，致命问题阻断"),
        ("b", "输出：结构化 dict（楼层 → 房间/构件，固定键名 → 规则引擎直接取值）"),
    ], s_size=S)
    notes(s, "这一页是读取侧的完整算法链。核心观点：全程确定性算法，没有一步靠大模型猜。打开文件后先解决单位——文件没声明就看楼层标高是 3 还是 3000 反推。然后分别用两种关系把楼层里的构件和房间捞出来。房间面积从几何算：曲面化、投影、凸包、鞋带公式。门宽走三级回退，门归房间优先边界关系、没有就几何兜底。疏散距离用图算法。结构构件截面和跨度一起提取。最后系统给模型打健康度分，致命问题直接叫停，不让烂模型进入审查。")


def s05(p):
    s = blank(p)
    title(s, "现有文件 + RAG 知识库搭建 + 数据↔资料比对", tag="资料端")
    block(s, Inches(0.6), Inches(1.55), Inches(12.1), [
        ("h", "① 现有文件结构"),
        ("s", "bim/samples/：FZK_Haus.ifc / office_3f.ifc / IfcOpenHouse_IFC4.ifc"),
        ("s", "data/standards/{cn,sg}/：civil / eee / building / structural_rules.json"),
        ("s", "data/rag_kb/：SG_Free / SG_Paid_Index / CN_Mandatory / CN_Reference"),
        ("s", "scripts/rag_kb/：kb_manifest.json / ocr_pdf.py / cleaning_rules.py / …"),
        ("blank",),
        ("h", "② RAG 知识库搭建步骤与方法"),
        ("s", "1. 搜集：锁定官方来源（SSO / gov.cn / BCA·SCDF），判定免费/付费 → kb_manifest.json（36 份）"),
        ("s", "2. 下载：curl 全文；证书吊销失败用 --ssl-no-revoke 绕过；付费仅索引"),
        ("s", "3. 转文本：文字层 PDF → pypdf 抽取；扫描版（GB 55037）→ RapidOCR → 36,921 字符"),
        ("s", "4. 落盘：按目录归档 data/rag_kb/"),
        ("s", "5. 清洗：cleaning_rules.py（窗口化 ±80 字符 + 词边界）→ 修正 FC23 6 处"),
        ("s", "6. 定位：agents/reviewer.py 拿 code_ref → 语料关键词匹配 → 证据片段"),
        ("blank",),
        ("h", "③ 数据端 ↔ 资料端比对 = 字段名契约"),
        ("s", "数据端 dict：door_width_m / depth_m / span_m …"),
        ("s", "规则 JSON：field / operator / threshold / condition_field / code_ref"),
        ("s", "比对：value OPERATOR threshold（8 操作符）；依据 = code_ref 标准名"),
    ], s_size=14)
    notes(s, "这一页把「有什么」和「怎么比」讲清楚。文件分两堆：数据端是样例模型和规则，资料端是 RAG 知识库。知识库搭建是一条完整流水线：官方来源搜集 → 判定免费付费 → 下载全文或只建索引 → 文字层直接抽、扫描版 OCR → 按目录落盘 → 清洗工具校验过期引用和错误条款号（实测修掉 FC23 6 处）→ 判定时拿规范引用去语料定位证据片段。数据怎么和资料比？一张「字段名契约」：数据端每个值有固定键名（如 door_width_m），规则端用同一个键名取值，引擎做值 vs 阈值的比对，规范依据用 code_ref 指向标准名。")


def s06(p):
    s = blank(p)
    title(s, "框架大更新：5 Agent 架构 + 四专业", tag="技术架构")
    block(s, Inches(0.6), Inches(1.5), Inches(12.1), [
        ("h", "为什么采用 Agent"),
        ("b", "单一模型职责过载、专业知识混装 → 按职责拆分：读/判/库/说/复核，可插拔、可扩展"),
    ], b_size=16)
    rows = [
        ("角色", "核心职责", "执行性质"),
        ("① BIM Reader", "读 IFC + 模型可用性 QA（健康度 0-100，致命阻断）", "确定性"),
        ("② Compliance Judge", "四专业规则判定 + 跨专业冲突检测", "确定性"),
        ("③ KB Steward", "规范知识库 RAG + 动态阈值解析", "检索确定+LLM 辅助"),
        ("④ Reasoner", "LLM 报告叙述 / 跨专业冲突协调", "LLM"),
        ("⑤ Independent Reviewer", "规范原文复核（概念保留，UI 已移除）", "确定性"),
    ]
    tbl = s.shapes.add_table(6, 3, Inches(0.6), Inches(2.35), Inches(6.4), Inches(3.4)).table
    tbl.columns[0].width = Inches(2.2)
    tbl.columns[1].width = Inches(7.4)
    tbl.columns[2].width = Inches(2.6)
    tbl.rows[0].height = Inches(0.4)
    for i in range(1, 6):
        tbl.rows[i].height = Inches(0.6)
    for r in range(6):
        for c in range(3):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE
            cell.margin_left = Inches(0.08)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            para(cell.text_frame, rows[r][c], 13, bold=(r == 0 or c == 0 or c == 2), first=True, space_after=0)
    block(s, Inches(0.6), Inches(5.9), Inches(12.1), [
        ("h", "四专业：文件四分类，引擎按对象类型跑"),
        ("b", "规则文件：{civil, eee, building, structural}_rules.json（消防 / 电气 / 全楼 / 结构）"),
        ("b", "引擎不挑文件：加载标准时四份一起读入，靠 field/condition_field 自动匹配"),
        ("b", "判定对象：civil/eee → 逐房间；structural → 逐构件；building → 整楼聚合"),
    ], b_size=15)
    notes(s, "架构上做了大重构：从「一个东西什么都干」拆成 5 个角色 Agent——读、判、库、说、复核各管一段。核心原则：读和判必须 100% 确定，同一模型任何时候跑结果一样；大模型只负责「说」——写报告、做解释，绝不让它拍板数值。关于四专业：规则文件按专业分成四类，但引擎不是先按专业挑规则——它把四份一起读进来，靠每条规则的 field 和 condition_field 自动匹配到对应对象：消防、电气逐房间判，结构逐构件判，全楼参数整楼判一次。分文件是维护方便，执行靠字段匹配。")


def s07(p):
    s = blank(p)
    title(s, "平面视图是如何得到的", tag="数据可视化")
    block(s, Inches(0.6), Inches(1.6), Inches(12.1), [
        ("h", "从 IFC 几何 → 平面图：算法实现"),
        ("b", "① 几何提取：ifcopenshell.geom.create_shape 把空间实体曲面化"),
        ("b", "② 投影与凸包：三维坐标 XY 投影 → Andrew 单调链算法构建房间凸包"),
        ("b", "③ 计算与渲染：鞋带公式算面积 → Pillow 渲染 → 3D 等轴 / 每层俯视 / 建筑剖面"),
        ("blank",),
        ("h", "技术优势与边界"),
        ("b", "数据驱动：模型改一笔，图即变；审查结果与平面图永远对齐"),
        ("b", "边界如实：凹形房间面积轻微高估（凸包特性），代码已标注"),
    ], b_size=17)
    notes(s, "大家看到的结果页平面图，不是手绘的，也不是从图纸里抠的——全是算法从模型几何里「长」出来的。三步：几何接口读出房间轮廓 → 投影 + 凸包简化 → 鞋带公式算面积 + Pillow 渲染，3D 等轴、每层俯视、剖面都能出。好处是数据驱动：模型改一笔，图跟着变，审查结果和平面图永远对得上。代价如实说：凸包对凹形房间面积会轻微高估，这是已知边界。")


def s08(p):
    s = blank(p)
    title(s, "遇到的问题：动态阈值（STR-002）", tag="真实工程问题")
    block(s, Inches(0.6), Inches(1.55), Inches(12.1), [
        ("h", "实际案例（FZK 主梁：实测高 0.24m · 跨度 4.08m）"),
        ("b", "传统死板规则 阈值=跨度/10：4.08/10≈0.41m → 0.24<0.41 → 误判违规"),
        ("b", "动态公式 max(0.2, min(跨度,6)/20)：max(0.2, 4.08/20)=0.204m → 0.24>0.204 → 通过"),
        ("blank",),
        ("h", "为什么要动态"),
        ("b", "① 几何数据 ≠ 设计意图：IFC 提取的是构件全长，不是设计跨度（檩条全长 13m、支座 4m）→ 几何噪音放大成假阳性"),
        ("b", "② 规范是「关系式」而非固定数：梁高本质「按跨高比确定」→ 公式反而更贴近规范逻辑"),
        ("blank",),
        ("h", "什么时候动态 / 什么时候静态"),
        ("b", "绝对量·构造不变值（柱 300mm、板 100mm、距水源 1.5m、坡道 1:12）→ 静态"),
        ("b", "随上下文变化的比例量（梁高按跨高比、疏散按路径、防火分区按面积、出口按人数）→ 动态"),
        ("b", "纪律：公式必须带构造底线 + 公开假设（0.2 是底线；6m 封顶、L/20 是住宅近似，已入决策文档）"),
    ], b_size=16)
    notes(s, "过程中遇到一个真实工程问题：规则写死会误判。案例是 FZK 的一根主梁，实测 0.24 米、跨度 4.08 米。死板规则「跨度的十分之一」给出阈值 0.41 米，0.24 小于 0.41，判违规——误判。动态公式按跨度算出 0.204 米，0.24 大于 0.204，通过——这才对。为什么必须动态？一是 IFC 的几何全长不等于设计跨度，死规则会被噪音带偏；二是规范本身常是「按跨高比」这种关系式，公式比死数更贴近规范逻辑。但我们守一条分界线：绝对量（构造底线）必须静态，比例量才动态；动态时必须带底线、公开假设。这不叫不按规范，是「按规范的关系式逻辑 + 对几何噪音做工程修正」。")


def s09(p):
    s = blank(p)
    title(s, "下一步（4 个方向）")
    block(s, Inches(0.6), Inches(1.6), Inches(12.1), [
        ("h", "① 完善 RAG —— 加入语义搜索"),
        ("b", "现状：关键词子串匹配，只能定位到章节/标题上下文；目标：embedding 向量检索，提高证据准确率"),
        ("h", "② 付费标准是否接入（版权权衡）"),
        ("b", "现状：付费标准仅索引，判定标「请查阅原文」；权衡：采购全文 / 结构化索引 / 维持人工核对"),
        ("h", "③ 字段定位 → 条款定位"),
        ("b", "现状：规则比对到「字段 vs 阈值」，依据只有标准名；目标：证据钉到「标准第 X 条原文」"),
        ("h", "④ 平面图优化"),
        ("b", "现状：凸包对凹形房间面积轻微高估；目标：多边形布尔运算/裁剪，修正拟合精度"),
    ], h_size=18, b_size=16)
    notes(s, "下一步四个方向。一是知识库加语义搜索，现在是关键词匹配、证据只能到章节；二是付费标准接入问题——涉及版权，我们倾向守住合规底线，具体是买全文、做索引还是继续人工，需要权衡；三是条款定位——现在只能把判定钉到字段和标准名，还钉不到规范第几条原文，这是可审计性的关键一步；四是平面图的几何优化。谢谢各位老师，欢迎提问。")


def main():
    p = Presentation()
    p.slide_width = SLIDE_W
    p.slide_height = SLIDE_H

    build = [
        ("封面", cover),
        ("上周两个问题 → 解决方案", s02),
        ("自生成 IFC + 步骤", s03),
        ("读取 IFC 算法链", s04),
        ("文件 + RAG + 比对", s05),
        ("5 Agent + 四专业", s06),
        ("平面视图如何得到", s07),
        ("动态阈值", s08),
        ("下一步 4 方向", s09),
    ]
    for i, (name, fn) in enumerate(build, 1):
        fn(p)
        print(f"  [{i:2d}/{len(build)}] {name}")

    out_dir = Path(__file__).parent.parent / "output"
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "ArchCheck_汇报PPT_2026-08-13_V2.pptx"
    p.save(out)
    print(f"\n已保存: {out}")
    print(f"共 {len(p.slides)} 页 · 纯白底 · 宋体 · 中文 · 讲解词在各页备注")


if __name__ == "__main__":
    main()
