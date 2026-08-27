# -*- coding: utf-8 -*-
"""Generate the roadmap-progress deck (pure white, large fonts, English).

Usage:  python scripts/make_ppt_roadmap.py
Output: presentation/2026-08-25/RoadmapProgress_2026-08-25.pptx
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "presentation" / "2026-08-25" / "RoadmapProgress_2026-08-25.pptx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
DARK = RGBColor(0x22, 0x22, 0x22)
GRAY = RGBColor(0x55, 0x55, 0x55)
RED = RGBColor(0xB0, 0x30, 0x30)
GREEN = RGBColor(0x1E, 0x6B, 0x3C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT = RGBColor(0xF2, 0xF4, 0xF7)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide():
    return prs.slides.add_slide(BLANK)


def _set_bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE


def tx(slide, left, top, width, height, lines, size=20, color=DARK,
       bold=False, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for text, level in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = text
        p.level = level
        p.alignment = align
        p.space_after = Pt(6)
        run = p.runs[0]
        run.font.size = Pt(size if level == 0 else max(size - 3, 12))
        run.font.bold = bold or (level == 0)
        run.font.color.rgb = color
        run.font.name = font
    return box


def title_bar(slide, text, size=30):
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(11.9), Inches(0.95))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.runs[0].font.size = Pt(size)
    p.runs[0].font.bold = True
    p.runs[0].font.color.rgb = NAVY
    p.runs[0].font.name = "Calibri"
    ln = slide.shapes.add_shape(1, Inches(0.7), Inches(1.28), Inches(11.9), Pt(2.2))
    ln.fill.solid()
    ln.fill.fore_color.rgb = NAVY
    ln.line.fill.background()
    ln.shadow.inherit = False
    return box


def make_table(slide, left, top, width, rows, col_widths, size=16,
               header_fill=NAVY, header_color=WHITE):
    n_rows, n_cols = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top),
                                   Inches(width), Inches(0.45 * n_rows))
    tbl = shape.table
    for j, w in enumerate(col_widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = Inches(0.06)
            cell.margin_right = Inches(0.06)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = 1
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = header_fill
            else:
                cell.fill.fore_color.rgb = WHITE if i % 2 else LIGHT
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            if not p.runs:
                continue
            run = p.runs[0]
            run.font.size = Pt(size if i else size)
            run.font.bold = (i == 0)
            run.font.color.rgb = header_color if i == 0 else DARK
            run.font.name = "Calibri"
    return shape


# ----------------------------------------------------------------------------
# Slide 1 · Title
# ----------------------------------------------------------------------------
s = add_slide()
_set_bg(s)
tx(s, 1.4, 2.0, 10.5, 1.4,
   [("The Road to Production-Ready Compliance Checking", 0)], size=40, color=NAVY, bold=True)
tx(s, 1.4, 3.3, 10.5, 0.9,
   [("A 6-tier industry-standard roadmap · progress & next steps", 0)], size=24, color=GRAY)
tx(s, 1.4, 5.7, 10.5, 0.6,
   [("ArchCheck · 2026-08-25", 0)], size=18, color=GRAY)

# ----------------------------------------------------------------------------
# Slide 2 · Roadmap overview
# ----------------------------------------------------------------------------
s = add_slide()
_set_bg(s)
title_bar(s, "The Roadmap at a Glance — 6 Tiers")
make_table(
    s, 0.9, 1.6, 11.6,
    [["Tier", "Done (this week)", "Plan (next)", "Progress"],
     ["1 · IFC reading", "spaces · doors · structure · MEP · electrical read; exact area / volume / material; mixed-unit fix",
      "connections · distribution systems · openings · quantity sets · profile shapes", "~70%"],
     ["2 · Quality validation", "QA 4 checks (units / spaces / geometry / boundaries); honest “not checked / area unknown”",
      "IDS parser + semantic validation (missing attributes, wrong classification)", "~40%"],
     ["3 · Geometry & quantities", "exact area (triangle projection) · exact volume · material take-off · basic clash pre-check",
      "boolean operations (CSG kernel) · swept solids · exact sections", "~85%"],
     ["4 · Clash & coordination", "MEP vs structure bbox + point-in-mesh confirmation (real vs bbox noise)",
      "full-discipline clash · clear-height checks · grading (hard / soft / clearance)", "early"],
     ["5 · Rules", "5 disciplines · 16 rules (CN + SG) · clause → rule tool",
      "16 → 50+ (fire / egress / MEP) · context & exceptions · cross-rule linkage", "~15%"],
     ["6 · Delivery & trust", "BCF 2.1 · COBie CSV · HTML/Excel/PNG · 4 samples run",
      "BCF viewpoints · fuller COBie · IFC with annotations · 10+ real projects · benchmark", "~35%"]],
    [2.6, 4.0, 3.7, 1.3], size=11)
tx(s, 0.9, 6.0, 11.6, 0.9,
   [("Foundation built: precise geometry + real samples + multi-discipline rules + delivery formats — now we build up.", 0)],
   size=20, color=NAVY, bold=True)

# ----------------------------------------------------------------------------
# Slide 3 · Breakthrough: precise geometry & quantities
# ----------------------------------------------------------------------------
s = add_slide()
_set_bg(s)
title_bar(s, "This Week's Breakthrough — Precise Geometry & Quantities (Tier ③)", size=27)
pts = [
    ("Exact area", "triangle projection replaces convex hull — L-shaped room 36 → 27 m², no over-estimate"),
    ("Exact volume", "signed-tetrahedron — office_3f slab 96.9 m³ · Schependomlaan structural total 287.4 m³"),
    ("Material & quantity take-off", "per-material volume + per-type counts → “structural quantities by material”"),
    ("Mixed-unit repair", "models declared in mm but modelled in metres auto-corrected (Schependomlaan)"),
    ("Basic clash pre-check", "MEP vs structure bbox + point-in-mesh confirmation — real penetration, not bbox noise"),
]
y = 1.6
for head, body in pts:
    tx(s, 0.9, y, 11.6, 0.95,
       [(head, 0), ("• " + body, 1)], size=19, color=DARK)
    y += 1.02
tx(s, 0.9, y + 0.12, 11.6, 0.7,
   [("Unlocks what's downstream: quantities, clash detection, clear-height checks — all depend on real geometry.", 0)],
   size=18, color=GREEN, bold=True)

# ----------------------------------------------------------------------------
# Slide 4 · Done vs still open
# ----------------------------------------------------------------------------
s = add_slide()
_set_bg(s)
title_bar(s, "Where We Stand — Done vs Still Open (honest)")
# left column: done
tx(s, 0.9, 1.55, 5.7, 0.55, [("Already in place", 0)], size=21, color=GREEN, bold=True)
done = [
    "5-discipline rule engine (CN + SG)",
    "BCF 2.1 / COBie export",
    "Clause → machine-rule tooling",
    "Semantic knowledge-base search",
    "Precise geometry, volumes, materials",
    "Real samples: Duplex_MEP, Schependomlaan",
]
y = 2.25
for d in done:
    tx(s, 0.9, y, 5.7, 0.6, [("• " + d, 0)], size=17, color=DARK)
    y += 0.62
# right column: open
tx(s, 6.9, 1.55, 5.7, 0.55, [("Still open (flagged, not hidden)", 0)], size=21, color=RED, bold=True)
opn = [
    "Boolean operations (openings) — needs CSG kernel",
    "IDS-based quality validation",
    "Rule volume: 16 → hundreds",
    "Full-discipline clash + clear-height checks",
    "BCF viewpoints · full COBie",
    "Scale benchmark vs commercial tools",
]
y = 2.25
for d in opn:
    tx(s, 6.9, y, 5.7, 0.6, [("• " + d, 0)], size=17, color=DARK)
    y += 0.62
tx(s, 0.9, 6.35, 11.6, 0.7,
   [("Everything we can't do yet is stated explicitly — no silent omissions.", 0)],
   size=18, color=GRAY, bold=True)

# ----------------------------------------------------------------------------
# Slide 5 · Next steps
# ----------------------------------------------------------------------------
s = add_slide()
_set_bg(s)
title_bar(s, "Next Steps — Prioritised")
rows = [
    ("Priority", "Action", "Effort"),
    ("1", "Rule expansion: 16 → 50+, fire / egress / MEP first", "2–3 wks"),
    ("2", "IDS quality checking (subset parser)", "1–2 wks"),
    ("3", "Clash deepening: full-discipline + clear-height + grading", "2 wks"),
    ("4", "Delivery: BCF viewpoints · fuller COBie · 10+ real projects", "1–2 wks"),
]
make_table(s, 0.9, 1.6, 11.6, rows, [1.4, 7.6, 2.6], size=16)
tx(s, 0.9, 4.05, 11.6, 0.5, [("Long-term hard items", 0)], size=20, color=NAVY, bold=True)
tx(s, 0.9, 4.6, 11.6, 1.0,
   [("• Boolean operations (CSG kernel)  ·  full rule coverage  ·  benchmark vs commercial tools", 0)],
   size=18, color=DARK)
tx(s, 0.9, 6.1, 11.6, 0.7,
   [("Foundation is built — next is rule volume, IDS validation, and clash depth.", 0)],
   size=20, color=GREEN, bold=True)

prs.save(OUT)
print("saved:", OUT)
